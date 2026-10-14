"""Tests for the Raw data -> dashboard_data.xlsx feeding."""
from __future__ import annotations

import os
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Alignment

from feeding import upload as upload_mod
from feeding.__main__ import main
from feeding.build import OUTPUT_NAME, collect, write_workbook
from feeding.closing import FS_COLUMNS, normalize_bs_label, parse_closing_pkg
from feeding.forecast import (latest_per_version, pick_latest, read_forecast_block,
                              read_forecast_history, version_of)
from feeding.pfe import latest_pfe_file, read_pfe_vendors
from feeding.pfe_pivot import build_project_pivot, read_project_scope, scoped_rows
from feeding.projects import parse_cost_analysis
from feeding.sources import data_dir, load_values, month_from_name, number
from feeding.vendor import latest_vendor_file, max_posting_date, read_vendor_names, read_vendor_table

REAL_BASE = Path(__file__).resolve().parent.parent
MONTHS = ["Jan", "Feb", "Mar"]


# ---------- builders for small synthetic workbooks ----------
def make_closing_pkg(path: Path, sales=(100.0, 200.0, 0.0)) -> Path:
    wb = openpyxl.Workbook()
    pl = wb.active
    pl.title = "1-4 P&L(COGS Components)"
    for i, m in enumerate(MONTHS):
        pl.cell(2, 5 + i, m)
    rows_b = [("Sales", sales), ("COGS", (80, 160, 0)), ("SG&A", (5, 6, 0)),
              ("Operating Income", (15, 34, 0)), ("EBITDA", (20, 40, 0)),
              ("Non-Operating Income", (1, 2, 0)), ("Non-Operating Expense", (3, 4, 0)),
              ("Net Income", (13, 32, 0))]
    r = 4
    for label, vals in rows_b:
        pl.cell(r, 2, label)
        for i, v in enumerate(vals):
            pl.cell(r, 5 + i, v)
        r += 1
    for label, vals in [("Material ", (50, 100, 0)), ("Labor", (20, 40, 0)), ("Overhead", (10, 20, 0)),
                        ("Submaterial ", (1, 1, 0))]:
        pl.cell(r, 3, label)
        for i, v in enumerate(vals):
            pl.cell(r, 5 + i, v)
        r += 1

    bs = wb.create_sheet("1-5. BS")
    for i, m in enumerate(MONTHS):
        bs.cell(7, 5 + i, m)
    bs_rows = [("SUM", "I. Current Assets", 60), ("SUM", "(2) Inventories", 25),
               (110100, "Cash", 10), (110500, "AR", 15),
               ("SUM", "II. Non-current assets held for sale", 0),
               ("SUM", "II. Non-current assets", 40), ("SUM", "Total Assets", 100),
               ("SUM", "I. Current liabilities", 30), (210200, "AP", 12), (210500, "STD", 8),
               ("SUM", "II. Non-current liabilities held for sale", 0),
               ("SUM", "III. Non-current liability", 20), (230300, "LTD", 18)]
    for r, (code, label, val) in enumerate(bs_rows, start=8):
        bs.cell(r, 2, code)
        bs.cell(r, 4, label)
        for i in range(3):
            bs.cell(r, 5 + i, val + i)
    wb.save(path)
    return path


def make_pivot(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PIVOT"
    rows = [("Margin Analysis 2026", None, None, None, None),
            ("Row Labels", "Qty", "Rev", "Cost", "PL"),
            ("VESTAS", 3, -300, -270, 30),
            ("  Project A", 1, -100, -90, 10),
            ("  Project B", 2, -200, -180, 20),
            ("GE", 1, -50, -60, -10),
            ("Project C", 1, -50, -60, -10),  # indented by style, not spaces
            ("Grand Total", 4, -350, -330, 20)]
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row, start=1):
            ws.cell(r, c, v)
    ws.cell(7, 1).alignment = Alignment(indent=1)
    wb.save(path)
    return path


def make_vendor(path: Path, with_icons=True) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    header = ["Account", "Document Number", "Document Type", "Posting Date", "Amount in local currency"]
    if with_icons:
        header = [None, None] + header
    ws.append(header)
    pad = [None, None] if with_icons else []
    ws.append(pad + [103270, 510001, "RE", datetime(2026, 8, 31), -10.0])
    ws.append(pad + [103270, 510002, "RE", datetime(2026, 7, 1), -5.0])
    ws.append([None] * len(header))
    wb.save(path)
    return path


def make_forecast(path: Path, bump: float = 0) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1.P&L_Summary"
    ws.cell(4, 2, "2026 Rolling")
    ws.cell(4, 15, "2+10")
    for i, m in enumerate(MONTHS):
        ws.cell(5, 3 + i, m)
    ws.cell(5, 15, "Total")
    totals = {"Sales Rev": 1000, "Gross Profit": 100, "O.I(Exc.HQ)": 60,
              "I.O(Inc.HQ)": 55, "Net Income(Exc.HQ)": 40}
    for r, (label, total) in enumerate(totals.items(), start=6):
        ws.cell(r, 2, label)
        ws.cell(r, 3, total / 10)
        ws.cell(r, 15, total + bump)
    wb.save(path)
    return path


@pytest.fixture
def raw_base(tmp_path: Path) -> Path:
    raw = tmp_path / "Raw data"
    for name in ("Financials", "Projects", "Forecast", "Vendor"):
        (raw / name).mkdir(parents=True)
    make_closing_pkg(raw / "Financials" / "CSWAM New Monthly Closing Pkg_2026.02.xlsx")
    make_pivot(raw / "Projects" / "Cost Analysis based on Project_2026.01.xlsx")
    make_pivot(raw / "Projects" / "Cost Analysis based on Project_2026.02.xlsx")
    make_forecast(raw / "Forecast" / "CSWAM Rolling FCST 1+11.xlsx")
    make_forecast(raw / "Forecast" / "CSWAM Rolling FCST 2+10.xlsx", bump=5)
    make_vendor(raw / "Vendor" / "fbl1n.xlsx")
    return tmp_path


# ---------- unit ----------
@pytest.mark.parametrize("label,expected", [
    ("I. Current Assets", "current assets"),
    ("Ⅰ. Current liabilities", "current liabilities"),
    ("(2) Inventories", "inventories"),
    ("Total Assets", "total assets"),
    ("II. Non-current assets", "non-current assets"),
])
def test_normalize_bs_label(label, expected):
    assert normalize_bs_label(label) == expected


def test_month_from_name():
    assert month_from_name("CSWAM New Monthly Closing Pkg_2026.08.xlsx") == "2026-08"
    assert month_from_name("CSWAM New Monthly Closing Pkg_2026.08_rev1.xlsx") == "2026-08"
    assert month_from_name("vendor.xlsx") is None


def test_load_values_reports_corrupt_file(tmp_path):
    bad = tmp_path / "Pkg_2026.01.xlsx"
    bad.write_bytes(b"not a zip")
    with pytest.raises(ValueError, match="파일 열기 실패"):
        load_values(bad)


def test_closing_pkg_without_jan_header_is_rejected(tmp_path):
    path = make_closing_pkg(tmp_path / "Pkg_2026.02.xlsx")
    wb = openpyxl.load_workbook(path)
    wb["1-4 P&L(COGS Components)"].cell(2, 5, "January")
    wb.save(path)
    with pytest.raises(ValueError, match="Jan"):
        parse_closing_pkg(path)


def test_text_posting_dates_are_parsed():
    table = [["Posting Date"], ["2026-08-31"], ["08/15/2026"], [None]]
    assert max_posting_date(table) == datetime(2026, 8, 31)


def test_upload_never_follows_redirects():
    handler = upload_mod._NoRedirect()
    assert handler.redirect_request(None, None, 302, "", {}, "https://elsewhere/") is None


def test_empty_vendor_folder_warns_about_fallback(raw_base):
    for f in (raw_base / "Raw data" / "Vendor").glob("*.xlsx"):
        f.unlink()
    make_vendor(raw_base / "vendor.xlsx", with_icons=False)
    res = collect(raw_base)
    assert any("Raw data/Vendor 비어있음" in w for w in res.warnings)


def test_number_ignores_text_and_bool():
    assert number("abc") == 0 and number(None) == 0 and number(True) == 0 and number(3) == 3.0


def test_pick_latest_forecast_uses_highest_n_and_rev_on_tie():
    names = ["CSWAM Rolling FCST 1+11.xlsx", "CSWAM Rolling FCST 10+2.xlsx",
             "CSWAM Rolling FCST 8+4.xlsx", "CSWAM Rolling FCST 10+2_rev1.xlsx"]
    assert pick_latest([Path(n) for n in names]).name == "CSWAM Rolling FCST 10+2_rev1.xlsx"
    assert pick_latest([Path("other.xlsx")]) is None
    assert version_of(Path("CSWAM Rolling FCST 8+4_rev1.xlsx")) == "8+4"


def test_parse_closing_pkg_maps_rows(tmp_path):
    rows = parse_closing_pkg(make_closing_pkg(tmp_path / "Pkg_2026.02.xlsx"))
    assert list(rows) == ["2026-01", "2026-02"]  # March has no sales
    jan = rows["2026-01"]
    assert jan["CurrentAssets"] == 60          # not the "Non-current assets" row
    assert jan["CurrentLiabilities"] == 30     # not "Non-current liability"
    assert jan["TotalLiabilities"] == 50 and jan["Equity"] == 50
    assert jan["MaterialCost"] == 50 and jan["LaborCost"] == 20 and jan["MfgOverhead"] == 10
    assert jan["EBITDA"] == 20 and jan["OperatingProfit"] == 15
    assert rows["2026-02"]["Cash"] == 11
    assert set(jan) == set(FS_COLUMNS)


def test_parse_closing_pkg_rejects_bad_name(tmp_path):
    with pytest.raises(ValueError):
        parse_closing_pkg(make_closing_pkg(tmp_path / "pkg.xlsx"))


def test_parse_cost_analysis(tmp_path):
    rows = parse_cost_analysis(make_pivot(tmp_path / "Cost Analysis based on Project_2026.08.xlsx"))
    assert [(r["Customer"], r["Project"]) for r in rows] == [
        ("VESTAS", "Project A"), ("VESTAS", "Project B"), ("GE", "Project C")]
    a = rows[0]
    assert a["Month"] == "2026-08" and a["Revenue"] == 100 and a["Cost"] == 90
    assert a["Margin"] == pytest.approx(0.1)


def test_read_vendor_table_drops_icon_columns_and_blank_rows(tmp_path):
    table = read_vendor_table(make_vendor(tmp_path / "v.xlsx"))
    assert table[0] == ["Account", "Document Number", "Document Type", "Posting Date",
                        "Amount in local currency"]
    assert len(table) == 3
    assert max_posting_date(table) == datetime(2026, 8, 31)


def test_read_vendor_table_requires_header(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(["foo", "bar"])
    wb.save(tmp_path / "bad.xlsx")
    with pytest.raises(ValueError):
        read_vendor_table(tmp_path / "bad.xlsx")


def test_latest_vendor_file_falls_back(tmp_path):
    fallback = make_vendor(tmp_path / "vendor.xlsx", with_icons=False)
    assert latest_vendor_file(None, fallback) == fallback
    assert latest_vendor_file(None, tmp_path / "missing.xlsx") is None


def test_data_dir_prefers_raw_data(raw_base):
    assert data_dir(raw_base, "Financials") == raw_base / "Raw data" / "Financials"
    assert data_dir(raw_base, "Nothing") is None


# ---------- integration (synthetic folder) ----------
def test_collect_and_write(raw_base):
    res = collect(raw_base)
    assert res.errors == []
    assert res.fs_months == ["2026-01", "2026-02"]
    assert res.sources["Forecast"] == "CSWAM Rolling FCST 2+10.xlsx"
    assert len(res.project_rows) == 6

    out = write_workbook(res, raw_base)
    out2 = write_workbook(res, raw_base)  # second write keeps a history copy
    assert out == out2 == raw_base / OUTPUT_NAME
    assert len(list((raw_base / "_history").glob("dashboard_data_*.xlsx"))) == 1

    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["FS", "Vendor", "VendorNames", "VendorSpend", "VendorCategories", "Project", "Forecast", "PFEVendors", "PFEProjects", "PFEProjectTable", "PipelineVsSAP", "ForecastHistory"]
    assert [c.value for c in wb["FS"][1]] == FS_COLUMNS
    assert wb["FS"]["A2"].value == "2026-01"
    assert wb["Forecast"]["B6"].value == "Sales Rev"


def test_collect_reports_missing_inputs(tmp_path):
    res = collect(tmp_path)
    assert len(res.errors) >= 4


def test_cli_check_only_does_not_write(raw_base, capsys):
    assert main(["--base", str(raw_base)]) == 0
    assert not (raw_base / OUTPUT_NAME).exists()
    assert "FS: 2개월" in capsys.readouterr().out


def test_cli_upload_without_secret_fails_cleanly(raw_base, monkeypatch):
    monkeypatch.delenv(upload_mod.SECRET_ENV, raising=False)
    assert main(["--base", str(raw_base), "--upload"]) == 2
    assert not (raw_base / OUTPUT_NAME).exists()  # nothing written when upload is impossible


def test_upload_sends_secret_header(raw_base, monkeypatch):
    sent = {}

    class FakeResp:
        status = 200
        def __enter__(self): return self
        def __exit__(self, *a): return False

    class FakeOpener:
        def open(self, req, timeout):
            sent.update(dict(req.header_items()))
            return FakeResp()

    monkeypatch.setenv(upload_mod.SECRET_ENV, "k")
    monkeypatch.setattr(upload_mod, "_OPENER", FakeOpener())
    f = raw_base / "x.bin"
    f.write_bytes(b"1")
    assert "HTTP 200" in upload_mod.upload(f, "x.bin")
    assert sent["X-secret"] == "k" and sent["X-filename"] == "x.bin"


# ---------- real Raw data (skipped when absent) ----------
@pytest.mark.skipif(data_dir(REAL_BASE, "Financials") is None, reason="no Raw data")
def test_real_raw_data_is_consistent():
    res = collect(REAL_BASE)
    assert res.errors == []
    for r in res.fs_rows:
        assert abs(r["TotalAssets"] - r["TotalLiabilities"] - r["Equity"]) < 1
        assert r["CurrentAssets"] < r["TotalAssets"]
        assert r["LaborCost"] > 0 and r["MaterialCost"] > r["COGS"] * 0.5
    block, path = read_forecast_block(data_dir(REAL_BASE, "Forecast"))
    assert len(block) == 33 and path.name == res.sources["Forecast"]


def test_cli_write_creates_output(raw_base, capsys):
    assert main(["--base", str(raw_base), "--write"]) == 0
    assert (raw_base / OUTPUT_NAME).exists()
    assert "생성 완료" in capsys.readouterr().out


def test_cli_stops_on_errors(tmp_path):
    assert main(["--base", str(tmp_path), "--write"]) == 1
    assert not (tmp_path / OUTPUT_NAME).exists()


def test_latest_per_version_keeps_one_file_per_version_in_order():
    names = ["CSWAM Rolling FCST 10+2.xlsx", "CSWAM Rolling FCST 2+10.xlsx",
             "CSWAM Rolling FCST 2+10_rev1.xlsx", "notes.xlsx"]
    assert [p.name for p in latest_per_version([Path(n) for n in names])] == [
        "CSWAM Rolling FCST 2+10_rev1.xlsx", "CSWAM Rolling FCST 10+2.xlsx"]


def test_read_forecast_history_collects_full_year_totals(raw_base):
    history, warns = read_forecast_history(raw_base / "Raw data" / "Forecast")
    assert warns == []
    assert [h["Version"] for h in history] == ["1+11", "2+10"]
    first, second = history
    assert first["ActualMonths"] == 1 and second["ActualMonths"] == 2
    assert first["GrossProfit"] == 100 and first["OperatingIncome"] == 60 and first["NetIncome"] == 40
    assert second["OperatingIncome"] == 65  # O.I(Exc.HQ), not the Inc.HQ row
    assert first["Sales"] == 1000


def test_read_forecast_history_warns_on_missing_rows(tmp_path):
    folder = tmp_path / "fc"
    folder.mkdir()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1.P&L_Summary"
    ws.cell(5, 3, "Jan")
    ws.cell(5, 15, "Total")
    ws.cell(6, 2, "Sales Rev")
    wb.save(folder / "CSWAM Rolling FCST 3+9.xlsx")
    history, warns = read_forecast_history(folder)
    assert history == [] and "rows not found" in warns[0]


def test_written_workbook_has_forecast_history_sheet(raw_base):
    out = write_workbook(collect(raw_base), raw_base)
    ws = openpyxl.load_workbook(out)["ForecastHistory"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("Version", "ActualMonths", "Sales", "GrossProfit", "OperatingIncome", "NetIncome")
    assert [r[0] for r in rows[1:]] == ["1+11", "2+10"]


def make_pfe_workbook(path: Path) -> Path:
    wb = openpyxl.Workbook()
    wb.active.title = "Other"
    ws = wb.create_sheet("PFE Vendor List")
    ws.append(["PFE", "Category", "Name", "Supplier"])
    ws.append(["Non-PFE", "Door", "Door Co", 103274])
    ws.append(["PFE", "Internal", "Bolt Co", 201570])
    ws.append(["PFE", "#N/A", "Tool Co", 103373])
    ws.append([None, None, None, None])
    wb.save(path)
    return path


def test_read_pfe_vendors_maps_columns_and_flags_missing_category(tmp_path):
    vendors, warns = read_pfe_vendors(make_pfe_workbook(tmp_path / "zmm009.xlsx"))
    assert vendors[0] == {"Supplier": "103274", "Name": "Door Co", "Category": "Door", "PFE": "Non-PFE"}
    assert [v["PFE"] for v in vendors] == ["Non-PFE", "PFE", "PFE"]
    assert vendors[2]["Category"] == ""  # Excel #N/A becomes blank
    assert len(warns) == 1 and "Tool Co" in warns[0]


def test_read_pfe_vendors_requires_sheet(tmp_path):
    wb = openpyxl.Workbook()
    wb.save(tmp_path / "x.xlsx")
    with pytest.raises(ValueError, match="PFE Vendor List"):
        read_pfe_vendors(tmp_path / "x.xlsx")


def test_collect_writes_pfe_vendors_when_folder_has_file(raw_base):
    (raw_base / "Raw data" / "PFE").mkdir()
    make_pfe_workbook(raw_base / "Raw data" / "PFE" / "zmm009.xlsx")
    assert latest_pfe_file(raw_base / "Raw data" / "PFE").name == "zmm009.xlsx"
    res = collect(raw_base)
    assert len(res.pfe_vendors) == 3 and "PFE 2" in res.sources["PFE"]
    ws = openpyxl.load_workbook(write_workbook(res, raw_base))["PFEVendors"]
    assert [c.value for c in ws[1]] == ["Supplier", "Name", "Category", "PFE"]
    assert ws.max_row == 4


def test_collect_warns_when_pfe_folder_missing(raw_base):
    res = collect(raw_base)
    assert res.pfe_vendors == [] and any("PFE 폴더" in w for w in res.warnings)


def make_invoice_workbook(path: Path) -> Path:
    wb = make_pfe_workbook(path) and openpyxl.load_workbook(path)
    y25 = wb.create_sheet("2025")
    y25.append(["Mat. Doc.", "Cust.Proj.No.", "Supplier", "Amount LC", "Pstng Date"])
    y25.append(["1", "ALPHA 10T", 103274, 300, "2025.07.07"])      # Non-PFE
    y25.append(["2", "ALPHA 10T", 201570, 100, "2025.08.01"])      # PFE
    y25.append(["3", "ALPHA 10T", 999999, 5000, "2025.08.01"])     # not on the list
    y26 = wb.create_sheet("2026")
    y26.append(["Material Document", "Customer Project No.", "Supplier", "Amt.in loc.cur.", "Posting Date"])
    y26.append(["4", "BETA 5T", "201570", "1,000.50", 46295])      # PFE, text amount, serial date
    y26.append(["5", "BETA 5T", "103274", "250.50-", "2026.01.02"])  # Non-PFE, SAP negative
    y26.append(["6", "GAMMA 1T", "201570", 40, "2026.02.01"])
    y26.append(["7", "GAMMA 1T", "103274", -40, "2026.02.01"])     # nets to zero → no ratio
    pivot = wb.create_sheet("PIVOT 2025 2026")
    pivot.append(["Cust.Proj.No.", "Name", None, None, None, None, None, None, "Project Name", "Total"])
    for name in ["GAMMA 1T", "ALPHA 10T", "OMEGA 9T", "Grand Total", "ALPHA 10T"]:
        pivot.append([None] * 8 + [name, 0])
    wb.save(path)
    return path


def test_build_project_pivot_splits_amounts_by_pfe_status(tmp_path):
    path = make_invoice_workbook(tmp_path / "zmm009.xlsx")
    vendors, _ = read_pfe_vendors(path)
    pv = build_project_pivot(path, {v["Supplier"]: v["PFE"] for v in vendors})
    assert pv.sheets == ["2025", "2026"]
    assert pv.lines_used == 6 and pv.lines_off_list == 1 and pv.lines_bad_amount == 0
    rows = {r["Project"]: r for r in pv.rows}
    assert rows["ALPHA 10T"] == {"Project": "ALPHA 10T", "NonPFE": 300, "PFE": 100, "Total": 400, "PFERatio": 0.25}
    beta = rows["BETA 5T"]
    assert beta["PFE"] == 1000.5 and beta["NonPFE"] == -250.5 and beta["Total"] == 750
    assert rows["GAMMA 1T"]["PFERatio"] is None          # total 0 → ratio not meaningful
    assert [r["Project"] for r in pv.rows][0] == "BETA 5T"  # sorted by total, largest first
    assert str(pv.first_date) == "2025-07-07" and str(pv.last_date) == "2026-09-30"


def test_build_project_pivot_rejects_sheet_without_amount_column(tmp_path):
    path = make_pfe_workbook(tmp_path / "z.xlsx")
    wb = openpyxl.load_workbook(path)
    wb.create_sheet("2026").append(["Customer Project No.", "Supplier"])
    wb.save(path)
    with pytest.raises(ValueError, match="Amount"):
        build_project_pivot(path, {"103274": "Non-PFE"})


def test_collect_writes_pfe_project_pivot(raw_base):
    (raw_base / "Raw data" / "PFE").mkdir()
    make_invoice_workbook(raw_base / "Raw data" / "PFE" / "zmm009.xlsx")
    res = collect(raw_base)
    assert [r["Project"] for r in res.pfe_projects] == ["BETA 5T", "ALPHA 10T", "GAMMA 1T"]
    ws = openpyxl.load_workbook(write_workbook(res, raw_base))["PFEProjects"]
    assert [c.value for c in ws[1]] == ["Project", "NonPFE", "PFE", "Total", "PFERatio"]
    assert ws.max_row == 4


def test_read_project_scope_uses_pivot_summary_order(tmp_path):
    scope, source = read_project_scope(make_invoice_workbook(tmp_path / "zmm009.xlsx"))
    assert scope == ["GAMMA 1T", "ALPHA 10T", "OMEGA 9T"]  # stops at Grand Total
    assert "PIVOT" in source


def test_read_project_scope_falls_back_to_relevant_sheet(tmp_path):
    path = make_pfe_workbook(tmp_path / "z.xlsx")
    wb = openpyxl.load_workbook(path)
    rel = wb.create_sheet("Relevant Projects Y25 to Y26")
    rel.append(["Mat. Doc.", "Cust.Proj.No."])
    for p in ["B 2T", "A 1T", "B 2T"]:
        rel.append(["1", p])
    wb.save(path)
    assert read_project_scope(path) == (["B 2T", "A 1T"], "Relevant Projects Y25 to Y26 시트")


def test_scoped_rows_keep_scope_order_and_flag_missing():
    rows, missing = scoped_rows(["B", "A", "Z"], [
        {"Project": "A", "NonPFE": 1, "PFE": 1, "Total": 2, "PFERatio": 0.5},
        {"Project": "B", "NonPFE": 3, "PFE": 1, "Total": 4, "PFERatio": 0.25},
        {"Project": "C", "NonPFE": 9, "PFE": 9, "Total": 18, "PFERatio": 0.5}])
    assert [r["Project"] for r in rows] == ["B", "A", "Z"] and missing == ["Z"]
    assert rows[2]["Total"] == 0 and rows[2]["PFERatio"] is None


def test_collect_builds_scoped_project_table(raw_base):
    folder = raw_base / "Raw data" / "PFE"
    folder.mkdir()
    make_invoice_workbook(folder / "ZMM009 Invoice Verification.xlsx")
    res = collect(raw_base)
    assert [r["Project"] for r in res.pfe_table] == ["GAMMA 1T", "ALPHA 10T", "OMEGA 9T"]
    assert res.pfe_table[1]["Total"] == 400 and res.pfe_table[1]["PFERatio"] == 0.25
    assert any("OMEGA 9T" in w for w in res.warnings)  # in scope, but no invoices
    ws = openpyxl.load_workbook(write_workbook(res, raw_base))["PFEProjectTable"]
    assert ws.max_row == 4


def test_read_vendor_names_from_vendor_list_sheet(tmp_path):
    path = make_vendor(tmp_path / "v.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb.create_sheet("Vendor list")
    ws.append(["Cty", "City", "Search Term", "Name", "Supplier"])
    ws.append(["US", "", "LEECO", "Leeco Steel, LLC", 103270])
    ws.append(["AM", "", "", "WARREN, C", "AME4201223"])
    ws.append(["US", "", "", "", 999])          # no name → skipped
    wb.save(path)
    assert read_vendor_names(path) == [{"Supplier": "103270", "Name": "Leeco Steel, LLC"},
                                       {"Supplier": "AME4201223", "Name": "WARREN, C"}]
    assert read_vendor_names(make_vendor(tmp_path / "plain.xlsx")) == []


def test_vendor_spend_nets_invoice_documents_and_picks_common_terms():
    from feeding.vendor_spend import vendor_spend
    table = [["Account", "Document Type", "Amount in local currency", "Terms of Payment"],
             ["103270", "RE", -100.0, "E115"],
             ["103270", "RE", -50.0, "E115"],
             ["103270", "KG", 20.0, "T030"],      # credit memo reduces spend
             ["103270", "AB", -999.0, "E115"],    # clearing document: not spend
             ["201203", "KR", -40.0, "E120"],
             ["AME4201223", "KR", -500.0, ""]]    # employee account: skipped
    rows = vendor_spend(table, {"103270": "Leeco"}, {"E115": "EOM + 115 days"})
    assert rows[0] == {"Supplier": "103270", "Name": "Leeco", "TermsCode": "E115",
                       "Terms": "EOM + 115 days", "Amount": 130.0}
    assert rows[1]["Supplier"] == "201203" and rows[1]["Terms"] == "E120" and len(rows) == 2


def test_read_payment_terms_sheet(tmp_path):
    from feeding.vendor_spend import read_payment_terms
    path = make_vendor(tmp_path / "v.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb.create_sheet("Payment Term")
    ws.append(["B000", "[Bank] Payment immediately"])
    ws.append(["E115", "[T/T] EOM + 115 days"])
    wb.save(path)
    assert read_payment_terms(path) == {"B000": "[Bank] Payment immediately", "E115": "[T/T] EOM + 115 days"}


def make_gl_workbook(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "GL Line item"
    ws.append(["Account", "Document Type", "Document Number", "Amount in local currency"])
    ws.append([620813, "KR", "D1", 90.0])      # service vendor: professional fee
    ws.append([420029, "KR", "D1", 10.0])
    ws.append([210410, "KR", "D1", -100.0])    # AP line (credit) ignored
    ws.append([990113, "RE", "D2", 98.0])      # material vendor: GR/IR main material
    ws.append([531390, "RE", "D2", 2.0])       # tiny supplies line must not win
    ws.append([450318, "RE", "D3", 50.0])      # OTA: price difference only
    names = wb.create_sheet("GL Name")
    names.append(["GL Account no", "GL Name"])
    for code, name in [(620813, "Professional fee-Service Fee"), (420029, "COGS-Other"),
                       (990113, "GR/IR Clearing-Main material"), (531390, "Supplies-Other"),
                       (450318, "Purchase price difference-MM"), (210410, "AP-Other")]:
        names.append([code, name])
    wb.save(path)
    return path


def test_vendor_categories_from_gl(tmp_path):
    from feeding.vendor_category import vendor_categories, vendor_documents
    table = [["Account", "Document Type", "Document Number"],
             ["103000", "KR", "D1"], ["103270", "RE", "D2"], ["105477", "RE", "D3"], ["103270", "AB", "D9"]]
    docs = vendor_documents(table)
    assert docs == {"D1": "103000", "D2": "103270", "D3": "105477"}  # clearing doc D9 not used
    cats = {c["Supplier"]: c for c in vendor_categories(make_gl_workbook(tmp_path / "GL Line item.xlsx"), docs)}
    assert cats["103000"]["CostCategory"] == "Professional fee-Service Fee"
    assert cats["103270"]["CostCategory"] == "" and cats["103270"]["MaterialCategory"] == "Main Material"
    assert cats["105477"]["CostCategory"] == "" and cats["105477"]["RuleCategory"] == "Logistics"


def test_vendor_file_pick_skips_gl_workbook(tmp_path):
    from feeding.vendor_category import find_gl_file
    make_vendor(tmp_path / "Vendor Line item.xlsx")
    make_gl_workbook(tmp_path / "GL Line item.xlsx")   # newer file, must not be taken as vendor data
    assert latest_vendor_file(tmp_path, tmp_path / "none.xlsx").name == "Vendor Line item.xlsx"
    assert find_gl_file(tmp_path).name == "GL Line item.xlsx"


def test_display_category_shortens_insurance():
    from feeding.vendor_category import display_category
    assert display_category("Insurance expense for prod. (health,GWL,Manulife)") == "Insurance"
    assert display_category("Professional fee-Service Fee") == "Professional fee-Service Fee"


def test_pipeline_project_matching():
    from feeding.pipeline_compare import match_sap_projects
    sap = ["ALLE CATT II 13T", "BUFFALO GAP 100T", "BUFFALO GAP 2T", "4Q25 114M G920", "HIGHLAND - 12A"]
    assert match_sap_projects("2026 Alle Cat II", sap) == ["ALLE CATT II 13T"]
    assert match_sap_projects("Buffalo Gap", sap) == ["BUFFALO GAP 100T", "BUFFALO GAP 2T"]
    assert match_sap_projects("2026 GE 114m Forecast", sap) == ["4Q25 114M G920"]   # explicit override
    assert match_sap_projects("2026 VESTAS Forecast", sap) == []                      # placeholder row
    assert match_sap_projects("2026 North Hills 15T", sap) == []


def make_cogs_package(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1-7. COGS Detail"
    ws.append([None, "PJT Name", "QTY", "Month", "Sales Total", "COGS Total", "Main Material"])
    ws.append([None, None, None, None, "Sales Revenue", "COGS", "Steel plate", "Flange", "Internal", "Aluminum"])
    ws.append([None, "BUFFALO GAP 100T", 4, "Jan", 800, 0, 400, 80, 60, 12])
    ws.append([None, "BUFFALO GAP 2T", 1, "Feb", 200, 0, 100, 20, 15, 3])
    ws.append([None, "NON-PROJECT", 0, "Jan", 50, 0, 0, 0, 0, 0])
    wb.save(path)
    return path


def make_mpp_pipeline(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MPP Invocing"
    ws.append([None])
    top = ["Customer", "Project Name"] + [None] * 15 + ["No of Sec (Act)", "Rev_(K USD) / Sec (Act)"]
    ws.append(top)
    sub = [None] * 20 + ["ASP_Plate", "Steel Handling", "ASP_Door Frame", "D/F Handling", "ASP_Flange",
                         "Flange_Handling"] + [None] * 5 + ["ASP_Internal", "Internal Hanling"]
    ws.append(sub)
    row = ["Vestas", "Buffalo Gap"] + [None] * 15 + [5, 0.21] + [None, 90, 10, 4, 0, 18, 0] + [None] * 5 + [12, 1]
    ws.append(row)
    ws.append(["Vestas", "2026 VESTAS Forecast"] + [None] * 15 + [5, 8])
    wb.save(path)
    return path


def test_pipeline_vs_sap_per_section(tmp_path):
    from feeding.pipeline_compare import compare, read_pipeline_unit_values, read_sap_project_costs
    sap = read_sap_project_costs(make_cogs_package(tmp_path / "Pkg_2026.08.xlsx"))
    assert sap["BUFFALO GAP 100T"]["QTY"] == 4 and sap["BUFFALO GAP 2T"]["Steel"] == 100
    pipe = read_pipeline_unit_values(make_mpp_pipeline(tmp_path / "Pipeline2026.xlsx"))
    assert pipe[0]["Revenue"] == 210 and pipe[0]["Steel"] == 100 and pipe[0]["Internal"] == 13
    rows, skipped = compare(pipe, sap)
    assert skipped == ["2026 VESTAS Forecast"]
    r = rows[0]
    assert r["SapProjects"] == "BUFFALO GAP 100T, BUFFALO GAP 2T" and r["SapSections"] == 5
    assert r["Revenue_SAP"] == 200 and r["Steel_SAP"] == 100 and r["DoorFrame_SAP"] == 3   # Aluminum
    assert r["DoorFrame_Pipeline"] == 4 and r["Flange_Pipeline"] == 18


def test_pipeline_reader_stops_at_first_total_row(tmp_path):
    from feeding.pipeline_compare import read_pipeline_unit_values
    path = make_mpp_pipeline(tmp_path / "Pipeline2026.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb["MPP Invocing"]
    ws.append([None, "Total"] + [None] * 15 + [0, 0])
    ws.append(["Customer", "Project Name"])                       # second table below
    ws.append(["Vestas", "Buffalo Gap"] + [None] * 15 + [5, 0.3])
    wb.save(path)
    assert [r["Project"] for r in read_pipeline_unit_values(path)] == ["Buffalo Gap", "2026 VESTAS Forecast"]


def test_pipeline_file_comes_from_sales_pipeline_folder(tmp_path):
    from feeding.pipeline_compare import PIPELINE_FOLDER, find_pipeline_file
    from feeding.sources import data_dir
    (tmp_path / "Raw data" / "PFE").mkdir(parents=True)
    make_mpp_pipeline(tmp_path / "Raw data" / "PFE" / "Old_Pipeline.xlsx")
    folder = tmp_path / "Raw data" / PIPELINE_FOLDER
    folder.mkdir()
    new = make_mpp_pipeline(folder / "CSAM_Pipeline2026_Oct.xlsx")
    assert find_pipeline_file([data_dir(tmp_path, PIPELINE_FOLDER)]) == new
