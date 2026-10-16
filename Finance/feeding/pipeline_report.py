"""Sales Pipeline vs SAP as a formatted Excel report (USD per section).

    python -m feeding.pipeline_report                       # newest pipeline in Raw data/Sales Pipeline
    python -m feeding.pipeline_report --pipeline <file>     # compare a specific pipeline file
    → writes Pipeline_vs_SAP_YYYY.MM.xlsx in the dashboard folder

Same numbers as the dashboard tab, plus Paint and a Material Total
(Steel + Flange + Internal + Door Frame + Paint). Δ, Δ% and Material Total are Excel formulas.
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .pipeline_compare import (MATERIALS, PIPELINE_FOLDER, compare, find_pipeline_file,
                               read_pipeline_unit_values, read_sap_project_costs)
from .sources import data_dir, list_xlsx

BASE_DIR = Path(__file__).resolve().parent.parent
GROUPS = [("Revenue", "Revenue"), ("Steel", "Steel"), ("Flange", "Flange"), ("Internal", "Internal"),
          ("DoorFrame", "Door Frame"), ("Paint", "Paint"), ("Material", "Material Total")]
SUB_HEADERS = ["Pipeline", "SAP", "Δ", "Δ%"]
INFO_COLUMNS = [("Customer", 14), ("Project", 26), ("SAP Projects", 34), ("Sec/Tower", 10), ("SAP Sections", 10)]
BOLD_SHARE = 0.10   # |Δ%| above this is bold, same as the dashboard
HEAD_ROW, FIRST_ROW = 4, 6
FIRST_GROUP_COL = len(INFO_COLUMNS) + 1

USD = "#,##0"
DELTA = "+#,##0;-#,##0;0"
SHARE = "+0.0%;-0.0%;0.0%"
HEAD_FILL = PatternFill("solid", fgColor="1F5C99")
SUB_FILL = PatternFill("solid", fgColor="E7F1FB")
THIN = Side(style="thin", color="D9E1EA")
GROUP_EDGE = Side(style="medium", color="9CBEE3")


def _col(group_index: int, offset: int) -> int:
    return FIRST_GROUP_COL + group_index * len(SUB_HEADERS) + offset


def _header(ws, last_col: int) -> None:
    for i, (name, width) in enumerate(INFO_COLUMNS, start=1):
        ws.merge_cells(start_row=HEAD_ROW, start_column=i, end_row=HEAD_ROW + 1, end_column=i)
        ws.cell(HEAD_ROW, i, name)
        ws.column_dimensions[get_column_letter(i)].width = width
    for g, (_, label) in enumerate(GROUPS):
        ws.merge_cells(start_row=HEAD_ROW, start_column=_col(g, 0), end_row=HEAD_ROW, end_column=_col(g, 3))
        ws.cell(HEAD_ROW, _col(g, 0), label)
        for k, sub in enumerate(SUB_HEADERS):
            ws.cell(HEAD_ROW + 1, _col(g, k), sub).fill = SUB_FILL
            ws.column_dimensions[get_column_letter(_col(g, k))].width = 9.5 if sub == "Δ%" else 10
    for r in (HEAD_ROW, HEAD_ROW + 1):
        for c in range(1, last_col + 1):
            cell = ws.cell(r, c)
            if r == HEAD_ROW or c <= len(INFO_COLUMNS):
                cell.fill = HEAD_FILL
                cell.font = Font(bold=True, color="FFFFFF")
            else:
                cell.font = Font(bold=True, color="1F5C99")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _row(ws, r: int, row: dict) -> None:
    info = [row["Customer"], row["Project"], row["SapProjects"] or "no actuals yet",
            row["PipelineSections"] or None, row["SapSections"] or None]
    for c, v in enumerate(info, start=1):
        ws.cell(r, c, v)
    for g, (metric, _) in enumerate(GROUPS):
        p, s = (get_column_letter(_col(g, k)) for k in (0, 1))
        if metric == "Material":   # sum of the material groups on this row
            for k in (0, 1):
                refs = ",".join(f"{get_column_letter(_col(i, k))}{r}"
                                for i, (m, _) in enumerate(GROUPS) if m in MATERIALS)
                ws.cell(r, _col(g, k), f'=IF(COUNT({refs})=0,"",SUM({refs}))')
        else:
            ws.cell(r, _col(g, 0), row[f"{metric}_Pipeline"])
            ws.cell(r, _col(g, 1), row[f"{metric}_SAP"])
        ws.cell(r, _col(g, 2), f'=IF(OR({p}{r}="",{s}{r}=""),"",{s}{r}-{p}{r})')
        ws.cell(r, _col(g, 3), f'=IF(OR({p}{r}="",{s}{r}="",{p}{r}=0),"",({s}{r}-{p}{r})/{p}{r})')
        for k, fmt in enumerate((USD, USD, DELTA, SHARE)):
            ws.cell(r, _col(g, k)).number_format = fmt
    ws.cell(r, 4).number_format = ws.cell(r, 5).number_format = USD


def _formats(ws, last_row: int, last_col: int) -> None:
    for r in range(FIRST_ROW, last_row + 1):
        for c in range(1, last_col + 1):
            cell = ws.cell(r, c)
            left = GROUP_EDGE if c >= FIRST_GROUP_COL and (c - FIRST_GROUP_COL) % len(SUB_HEADERS) == 0 else THIN
            cell.border = Border(left=left, right=THIN, top=THIN, bottom=THIN)
            cell.alignment = Alignment(horizontal="left" if c <= 3 else "right", vertical="center")
        ws.cell(r, 2).font = Font(bold=True)
        ws.cell(r, 3).font = Font(color="808080", size=9)
    for g, (metric, _) in enumerate(GROUPS):
        p, d, pct = (get_column_letter(_col(g, k)) for k in (0, 2, 3))
        delta_rng = f"{d}{FIRST_ROW}:{pct}{last_row}"
        # revenue above plan is favorable; costs below plan are favorable
        good, bad = (">", "<") if metric == "Revenue" else ("<", ">")
        ws.conditional_formatting.add(delta_rng, FormulaRule(
            formula=[f'AND(ISNUMBER(${d}{FIRST_ROW}),${d}{FIRST_ROW}{good}0)'], font=Font(color="2D8F65")))
        ws.conditional_formatting.add(delta_rng, FormulaRule(
            formula=[f'AND(ISNUMBER(${d}{FIRST_ROW}),${d}{FIRST_ROW}{bad}0)'], font=Font(color="D64545")))
        ws.conditional_formatting.add(f"{p}{FIRST_ROW}:{pct}{last_row}", FormulaRule(
            formula=[f'AND(ISNUMBER(${pct}{FIRST_ROW}),ABS(${pct}{FIRST_ROW})>{BOLD_SHARE})'], font=Font(bold=True)))


def write_report(rows: list[dict], out: Path, note: str) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pipeline vs SAP"
    last_col = _col(len(GROUPS) - 1, len(SUB_HEADERS) - 1)
    ws["A1"] = "Per-Section Revenue & Material Cost — Pipeline vs SAP (USD per section)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = note
    ws["A2"].font = Font(color="808080", size=9)
    _header(ws, last_col)
    for i, row in enumerate(rows):
        _row(ws, FIRST_ROW + i, row)
    last_row = FIRST_ROW + len(rows) - 1
    _formats(ws, last_row, last_col)
    ws.freeze_panes = ws.cell(FIRST_ROW, FIRST_GROUP_COL)
    ws.sheet_view.zoomScale = 85
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(out)
    return out


def build(base: Path = BASE_DIR, pipeline: Path | None = None) -> Path:
    fin = data_dir(base, "Financials")
    if pipeline is not None and not pipeline.is_file():
        raise FileNotFoundError(f"Pipeline 파일 없음: {pipeline}")
    pipeline = pipeline or find_pipeline_file([data_dir(base, PIPELINE_FOLDER)])
    if fin is None or pipeline is None:
        raise FileNotFoundError(f"Raw data/{PIPELINE_FOLDER} Pipeline 파일 또는 Raw data/Financials Closing Pkg 없음")
    package = list_xlsx(fin)[-1]
    rows, _ = compare(read_pipeline_unit_values(pipeline), read_sap_project_costs(package))
    month = re.search(r"(\d{4}\.\d{2})", package.name)
    note = (f"Pipeline: {pipeline.name} (MPP Invocing) · SAP: {package.name} (1-7 sheet, project total ÷ QTY) · "
            "Δ = SAP − Pipeline (green = favorable) · |Δ%| > 10% bold · Door Frame = SAP Aluminum · "
            "Paint = Paint + Zinc Wire + handling · Material Total = Steel + Flange + Internal + Door Frame + Paint")
    out = base / f"Pipeline_vs_SAP_{month.group(1) if month else 'latest'}.xlsx"
    return write_report(rows, out, note)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(prog="feeding.pipeline_report")
    parser.add_argument("--pipeline", type=Path, help="pipeline file (default: newest in Raw data/Sales Pipeline)")
    args = parser.parse_args()
    try:
        print(f"{build(pipeline=args.pipeline)} 생성 완료")
    except (OSError, ValueError) as exc:
        print(f"Pipeline vs SAP 엑셀 생성 실패: {exc}", file=sys.stderr)
        sys.exit(1)
