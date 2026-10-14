"""Assemble dashboard_data.xlsx from Raw data and run sanity checks."""
from __future__ import annotations

import os
import shutil
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import openpyxl

from .closing import FS_COLUMNS, build_fs_rows
from .forecast import HISTORY_COLUMNS, read_forecast_block, read_forecast_history, version_of
from .pfe import PFE_COLUMNS, latest_pfe_file, read_pfe_vendors
from .pfe_pivot import PROJECT_PIVOT_COLUMNS, build_project_pivot, read_project_scope, scoped_rows
from .projects import PROJECT_COLUMNS, build_project_rows
from .sources import data_dir, list_xlsx, month_from_name
from .vendor import latest_vendor_file, max_posting_date, read_vendor_names, read_vendor_table
from .vendor_spend import SPEND_COLUMNS, read_payment_terms, vendor_spend
from .vendor_category import CATEGORY_COLUMNS, find_gl_file, vendor_categories, vendor_documents
from .pipeline_compare import (COMPARE_COLUMNS, PIPELINE_FOLDER, compare, find_pipeline_file,
                               read_pipeline_unit_values, read_sap_project_costs)

OUTPUT_NAME = "dashboard_data.xlsx"
HISTORY_DIR = "_history"
COGS_TOLERANCE = 0.10  # Material+Labor+Overhead vs COGS (Scrap/Adj rows make up the rest)


@dataclass
class FeedResult:
    fs_rows: list[dict] = field(default_factory=list)
    vendor_table: list[list] = field(default_factory=list)
    vendor_names: list[dict] = field(default_factory=list)  # Supplier code → name
    vendor_spend: list[dict] = field(default_factory=list)  # total spend per vendor, largest first
    vendor_categories: list[dict] = field(default_factory=list)  # category per vendor from the GL
    project_rows: list[dict] = field(default_factory=list)
    forecast_block: list[list] = field(default_factory=list)
    forecast_history: list[dict] = field(default_factory=list)
    pfe_vendors: list[dict] = field(default_factory=list)
    pfe_projects: list[dict] = field(default_factory=list)
    pfe_table: list[dict] = field(default_factory=list)  # AMPC table: scoped projects, display order
    pipeline_compare: list[dict] = field(default_factory=list)  # Sales Pipeline vs SAP per section
    sources: dict[str, str] = field(default_factory=dict)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def fs_months(self) -> list[str]:
        return [r["Month"] for r in self.fs_rows]


def collect(base: Path) -> FeedResult:
    res = FeedResult()
    _collect_fs(base, res)
    _collect_vendor(base, res)
    _collect_projects(base, res)
    _collect_forecast(base, res)
    _collect_pfe(base, res)
    _collect_pipeline_compare(base, res)
    _check(res)
    return res


def _collect_fs(base: Path, res: FeedResult) -> None:
    folder = data_dir(base, "Financials")
    if folder is None:
        res.errors.append("Closing Package 폴더(Raw data/Financials)에 파일 없음")
        return
    res.fs_rows, warns = build_fs_rows(folder)
    res.warnings += warns
    files = list_xlsx(folder)
    res.sources["Financials"] = f"{folder} ({len(files)} files, 최신 {files[-1].name})"
    latest_pkg = month_from_name(files[-1].name)
    if res.fs_months and latest_pkg and res.fs_months[-1] != latest_pkg:
        res.warnings.append(
            f"최신 Closing Pkg는 {latest_pkg}인데 FS 마지막 월은 {res.fs_months[-1]} (해당 월 Sales 비어있음?)")


def _collect_vendor(base: Path, res: FeedResult) -> None:
    vendor_dir = data_dir(base, "Vendor")
    path = latest_vendor_file(vendor_dir, base / "vendor.xlsx")
    if path is None:
        res.errors.append("Vendor 파일 없음 (Raw data/Vendor 또는 vendor.xlsx)")
        return
    try:
        res.vendor_table = read_vendor_table(path)
    except ValueError as exc:
        res.errors.append(str(exc))
        return
    res.sources["Vendor"] = str(path)
    res.vendor_names = read_vendor_names(path)
    res.vendor_spend = vendor_spend(res.vendor_table, {v["Supplier"]: v["Name"] for v in res.vendor_names},
                                    read_payment_terms(path))
    if res.vendor_names:
        res.sources["Vendor 이름 목록"] = f"{len(res.vendor_names)}개 (Vendor list 시트)"
    _collect_vendor_categories(vendor_dir, res)
    if vendor_dir is None:
        res.warnings.append("Raw data/Vendor 비어있음 — 예전 vendor.xlsx 사용 중 (이번 달 FBL1N 넣었는지 확인)")
    latest = max_posting_date(res.vendor_table)
    if latest is None:
        res.warnings.append("Vendor 전기일(Posting Date)을 읽지 못함 — 기준일 확인 불가")
    res.sources["Vendor 최종 전기일"] = latest.strftime("%Y-%m-%d") if latest else "-"
    if latest and res.fs_months and latest.strftime("%Y-%m") < res.fs_months[-1]:
        res.warnings.append(
            f"Vendor 최종 전기일 {latest:%Y-%m-%d}가 FS 최신 월 {res.fs_months[-1]}보다 이전 — FBL1N 새로 받았는지 확인")


def _collect_vendor_categories(vendor_dir: Path | None, res: FeedResult) -> None:
    """Vendor categories from the GL line items (Document Number join); optional."""
    gl_path = find_gl_file(vendor_dir)
    if gl_path is None:
        res.warnings.append("GL Line item 파일 없음 — Vendor 분류는 알려진 업체만 표시")
        return
    try:
        res.vendor_categories = vendor_categories(gl_path, vendor_documents(res.vendor_table))
    except ValueError as exc:
        res.warnings.append(str(exc))
        return
    found = sum(1 for c in res.vendor_categories if c["CostCategory"] or c["RuleCategory"])
    res.sources["Vendor 분류"] = f"{gl_path.name} | 비용계정·규칙으로 분류 {found}/{len(res.vendor_categories)}개"


def _collect_projects(base: Path, res: FeedResult) -> None:
    folder = data_dir(base, "Projects")
    if folder is None:
        res.errors.append("Cost Analysis 폴더(Raw data/Projects)에 파일 없음")
        return
    res.project_rows, warns = build_project_rows(folder)
    res.warnings += warns
    res.sources["Projects"] = f"{folder} ({len(list_xlsx(folder))} files)"


def _collect_forecast(base: Path, res: FeedResult) -> None:
    folder = data_dir(base, "Forecast")
    if folder is None:
        res.errors.append("Forecast 폴더(Raw data/Forecast)에 파일 없음")
        return
    try:
        res.forecast_block, path = read_forecast_block(folder)
    except ValueError as exc:
        res.errors.append(str(exc))
        return
    res.sources["Forecast"] = path.name
    res.forecast_history, warns = read_forecast_history(folder)
    res.warnings += warns
    res.sources["Forecast 버전 이력"] = ", ".join(h["Version"] for h in res.forecast_history) or "-"
    actual_months = int(version_of(path).split("+")[0] or 0)
    if res.fs_months and actual_months < int(res.fs_months[-1][5:]):
        res.warnings.append(
            f"Forecast {version_of(path)} 는 실적 {actual_months}개월 기준 — FS는 {res.fs_months[-1]}까지 있음 (새 FCST 없으면 무시)")


def _collect_pfe(base: Path, res: FeedResult) -> None:
    """Optional: AMPC tab shows the PFE vendor list when Raw data/PFE has a workbook."""
    folder = data_dir(base, "PFE")
    path = latest_pfe_file(folder) if folder else None
    if path is None:
        res.warnings.append("PFE 폴더(Raw data/PFE)에 파일 없음 — AMPC 탭 벤더 목록 비어 있음")
        return
    try:
        res.pfe_vendors, warns = read_pfe_vendors(path)
    except ValueError as exc:
        res.warnings.append(str(exc))
        return
    res.warnings += warns
    pfe_count = sum(1 for v in res.pfe_vendors if v["PFE"] == "PFE")
    res.sources["PFE"] = f"{path.name} (벤더 {len(res.pfe_vendors)}개, PFE {pfe_count}개)"
    try:
        pivot = build_project_pivot(path, {v["Supplier"]: v["PFE"] for v in res.pfe_vendors})
    except ValueError as exc:
        res.warnings.append(f"PFE 프로젝트 피벗 실패: {exc}")
        return
    res.pfe_projects = pivot.rows
    period = f"{pivot.first_date} ~ {pivot.last_date}" if pivot.first_date else "-"
    res.sources["PFE 프로젝트 피벗"] = (
        f"시트 {', '.join(pivot.sheets)} | 기간 {period} | 프로젝트 {len(pivot.rows)}개 | "
        f"전표 {pivot.lines_used}줄 (목록 외 업체 {pivot.lines_off_list}줄 제외)")
    if pivot.lines_bad_amount:
        res.warnings.append(f"PFE 피벗: 금액을 읽지 못한 전표 {pivot.lines_bad_amount}줄 제외")
    scope, scope_source = read_project_scope(path)
    if not scope:
        res.warnings.append("PFE 대상 프로젝트 목록(PIVOT / Relevant Projects 시트) 없음 — AMPC 표 비어 있음")
        return
    res.pfe_table, missing = scoped_rows(scope, pivot.rows)
    res.sources["PFE 대상 프로젝트"] = f"{scope_source} {len(scope)}개"
    if missing:
        res.warnings.append(f"PFE 대상 프로젝트 중 2025/2026 시트에 전표 없음: {', '.join(missing)}")


def _collect_pipeline_compare(base: Path, res: FeedResult) -> None:
    """Optional: per-section Pipeline (MPP sheet) vs SAP (Closing Pkg 1-7) by project."""
    fin = data_dir(base, "Financials")
    pipeline = find_pipeline_file([data_dir(base, PIPELINE_FOLDER)])
    if fin is None or pipeline is None:
        res.warnings.append(f"Raw data/{PIPELINE_FOLDER} 에 Pipeline 파일 또는 Closing Pkg 없음 — Pipeline vs SAP 표 비어 있음")
        return
    package = list_xlsx(fin)[-1]   # newest package holds every month of the year
    try:
        res.pipeline_compare, skipped = compare(read_pipeline_unit_values(pipeline), read_sap_project_costs(package))
    except ValueError as exc:
        res.warnings.append(f"Pipeline vs SAP 실패: {exc}")
        return
    unmatched = [r["Project"] for r in res.pipeline_compare if not r["SapProjects"]]
    res.sources["Pipeline vs SAP"] = (f"{pipeline.name} ↔ {package.name} (1-7 시트) | 프로젝트 {len(res.pipeline_compare)}개"
                                      + (f" | 제외(Forecast 줄): {', '.join(skipped)}" if skipped else ""))
    if unmatched:
        res.warnings.append(f"Pipeline vs SAP: SAP 실적 없는 프로젝트 {', '.join(unmatched)}")


def _check(res: FeedResult) -> None:
    if not res.fs_rows:
        res.errors.append("FS 월 데이터 0건")
    if not res.project_rows:
        res.errors.append("Project 데이터 0건")
    if res.vendor_table and len(res.vendor_table) < 2:
        res.errors.append("Vendor 데이터 0건")

    for r in res.fs_rows:
        parts = r["MaterialCost"] + r["LaborCost"] + r["MfgOverhead"]
        if r["COGS"] and abs(parts - r["COGS"]) / abs(r["COGS"]) > COGS_TOLERANCE:
            res.warnings.append(f"{r['Month']}: Material+Labor+Overhead가 COGS와 10% 이상 차이 (원본 Adj 행 확인)")

    project_months = {r["Month"] for r in res.project_rows}
    missing = [m for m in res.fs_months if m not in project_months]
    if missing:
        res.warnings.append(f"Project 데이터 없는 월: {', '.join(missing)}")


def write_workbook(res: FeedResult, base: Path) -> Path:
    """Write dashboard_data.xlsx (previous version copied to _history/)."""
    out = base / OUTPUT_NAME
    wb = openpyxl.Workbook()
    fs = wb.active
    fs.title = "FS"
    fs.append(FS_COLUMNS)
    for r in res.fs_rows:
        fs.append([r[c] for c in FS_COLUMNS])

    vendor = wb.create_sheet("Vendor")
    for row in res.vendor_table:
        vendor.append(row)

    names = wb.create_sheet("VendorNames")
    names.append(["Supplier", "Name"])
    for v in res.vendor_names:
        names.append([v["Supplier"], v["Name"]])

    spend = wb.create_sheet("VendorSpend")
    spend.append(SPEND_COLUMNS)
    for r in res.vendor_spend:
        spend.append([r[c] for c in SPEND_COLUMNS])

    cats = wb.create_sheet("VendorCategories")
    cats.append(CATEGORY_COLUMNS)
    for r in res.vendor_categories:
        cats.append([r[c] for c in CATEGORY_COLUMNS])

    project = wb.create_sheet("Project")
    project.append(PROJECT_COLUMNS)
    for r in res.project_rows:
        project.append([r[c] for c in PROJECT_COLUMNS])

    forecast = wb.create_sheet("Forecast")
    for ri, row in enumerate(res.forecast_block, start=1):
        for ci, value in enumerate(row, start=1):
            if value is not None:  # keep the sheet range where the data starts
                forecast.cell(ri, ci, value)

    pfe = wb.create_sheet("PFEVendors")
    pfe.append(PFE_COLUMNS)
    for v in res.pfe_vendors:
        pfe.append([v[c] for c in PFE_COLUMNS])

    pfe_projects = wb.create_sheet("PFEProjects")
    pfe_projects.append(PROJECT_PIVOT_COLUMNS)
    for r in res.pfe_projects:
        pfe_projects.append([r[c] for c in PROJECT_PIVOT_COLUMNS])

    table = wb.create_sheet("PFEProjectTable")
    table.append(PROJECT_PIVOT_COLUMNS)
    for r in res.pfe_table:
        table.append([r[c] for c in PROJECT_PIVOT_COLUMNS])

    cmp_sheet = wb.create_sheet("PipelineVsSAP")
    cmp_sheet.append(COMPARE_COLUMNS)
    for r in res.pipeline_compare:
        cmp_sheet.append([r[c] for c in COMPARE_COLUMNS])

    history = wb.create_sheet("ForecastHistory")
    history.append(HISTORY_COLUMNS)
    for h in res.forecast_history:
        history.append([h[c] for c in HISTORY_COLUMNS])

    if out.exists():
        history = base / HISTORY_DIR
        history.mkdir(exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
        shutil.copy2(out, history / f"dashboard_data_{stamp}.xlsx")
    tmp = out.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    os.replace(tmp, out)
    return out
